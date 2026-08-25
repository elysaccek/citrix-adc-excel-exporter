from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
import re

def clean_sheet_name(name):
    clean_name = re.sub(r'[\\*?:/\[\]]', '', name)
    return clean_name[:31]

def create_excel_report(report_data, fields, excel_path, empty_row_template, company_name, datasets_data=None, policy_data=None):
    wb = Workbook()
    
    ws_main = wb.active
    ws_main.title = clean_sheet_name(company_name)

    CS_PALETTE = ["DDEBF7", "E2EFDA", "FFF2CC", "FCE4D6", "E6E6FA"] 
    VS_PALETTE = ["FFFFFF", "F2F2F2"]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(color="9C0006")
    
    thin_border = Border(left=Side(style='thin', color='BFBFBF'),
                         right=Side(style='thin', color='BFBFBF'),
                         top=Side(style='thin', color='BFBFBF'),
                         bottom=Side(style='thin', color='BFBFBF'))

    ws_main.append(fields)
    for cell in ws_main[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        cell.border = thin_border

    current_row = 2
    empty_rows_indices = []
    last_cs = None
    last_vs = None
    cs_idx = -1
    vs_idx = -1

    for row_data in report_data:
        if row_data == empty_row_template:
            empty_rows_indices.append(current_row)
            ws_main.append([""] * len(fields))
            for col_idx in range(1, len(fields) + 1):
                cell = ws_main.cell(row=current_row, column=col_idx)
                cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            current_row += 1
            continue

        current_cs = row_data.get("Content Switch", "")
        current_vs = row_data.get("Load Balancer", "")

        if current_cs != last_cs:
            cs_idx += 1
            last_cs = current_cs
        
        if current_vs != last_vs:
            vs_idx += 1
            last_vs = current_vs

        cs_color = CS_PALETTE[cs_idx % len(CS_PALETTE)]
        vs_color = VS_PALETTE[vs_idx % len(VS_PALETTE)]

        cs_fill = PatternFill(start_color=cs_color, end_color=cs_color, fill_type="solid")
        vs_fill = PatternFill(start_color=vs_color, end_color=vs_color, fill_type="solid")

        row_list = [row_data.get(f, "") for f in fields]
        ws_main.append(row_list)
        
        hl_lb_red = row_data.get("_highlight_lb_red", False)
        hl_az_red = row_data.get("_highlight_azure_red", False)
        hl_cs_missing = row_data.get("_highlight_cs_red", False)
        hl_lb_missing = row_data.get("_highlight_lb_missing_red", False)
        
        for col_idx in range(1, len(fields) + 1):
            cell = ws_main.cell(row=current_row, column=col_idx)
            cell.alignment = center_alignment
            cell.border = thin_border
            
            if col_idx <= 5:
                cell.fill = cs_fill
            else:
                cell.fill = vs_fill

            if hl_cs_missing:
                cell.fill = red_fill
                cell.font = red_font

            if hl_lb_missing and col_idx >= 9:
                cell.fill = red_fill
                cell.font = red_font
            
            if hl_lb_red and 13 <= col_idx <= 16:
                cell.fill = red_fill
                cell.font = red_font
                
            if hl_az_red and 17 <= col_idx <= 18:
                cell.fill = red_fill
                cell.font = red_font
                
        current_row += 1

    for row_idx in empty_rows_indices:
        ws_main.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(fields))

    last_row = ws_main.max_row

    start_cs_row = 2
    for row in range(2, last_row + 2):
        current_cs = ws_main.cell(row=row, column=1).value if row <= last_row else None
        prev_cs = ws_main.cell(row=start_cs_row, column=1).value
        
        if current_cs != prev_cs or row > last_row:
            if row - 1 > start_cs_row and prev_cs and prev_cs != "":
                for col_idx in range(1, 6): 
                    ws_main.merge_cells(start_row=start_cs_row, start_column=col_idx, end_row=row-1, end_column=col_idx)
            start_cs_row = row

    start_lb_row = 2
    for row in range(2, last_row + 2):
        current_lb = ws_main.cell(row=row, column=9).value if row <= last_row else None
        prev_lb = ws_main.cell(row=start_lb_row, column=9).value
        
        if current_lb != prev_lb or row > last_row:
            if row - 1 > start_lb_row and prev_lb and prev_lb != "":
                for col_idx in range(6, 13): 
                    ws_main.merge_cells(start_row=start_lb_row, start_column=col_idx, end_row=row-1, end_column=col_idx)
            start_lb_row = row

    for col in ws_main.columns:
        max_length = 0
        column_letter = col[0].column_letter
        for cell in col:
            try:
                val = str(cell.value) if cell.value is not None else ""
                lines = val.split("\n")
                for line in lines:
                    if len(line) > max_length:
                        max_length = len(line)
            except:
                pass
        adjusted_width = min(max(max_length + 6, 15), 70)
        ws_main.column_dimensions[column_letter].width = adjusted_width

    if policy_data:
        ws_pol = wb.create_sheet(title="POLICIES", index=1)
        
        pol_headers = [
            "Content Switch / Load Balancer", "CS IP / VS IP", "CS Port / VS Port", "CS Protokol / VS Protokol",
            "Policy", "Policy Type", "Policy Priority", "Policy Name", "Expression", 
            "Action Name", "Action Type", "Action diğer detayları", "Goto Expression", "Invoke Label Type"
        ]
        
        ws_pol.append(pol_headers)
        
        pol_header_fill = PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid")
        for cell in ws_pol[1]:
            cell.font = header_font
            cell.fill = pol_header_fill
            cell.alignment = center_alignment
            cell.border = thin_border
            
        POL_PALETTE = ["FFFFFF", "F2F2F2"]
        pol_idx = -1
        last_pol_vs = None
        
        current_row = 2
        for pol_row in policy_data:
            vs_name = pol_row.get("Content Switch / Load Balancer", "")
            if vs_name != last_pol_vs:
                pol_idx += 1
                last_pol_vs = vs_name
                
            bg_color = POL_PALETTE[pol_idx % len(POL_PALETTE)]
            fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
            
            row_values = [pol_row.get(h, "") for h in pol_headers]
            ws_pol.append(row_values)
            
            for col_idx in range(1, len(pol_headers) + 1):
                cell = ws_pol.cell(row=current_row, column=col_idx)
                cell.fill = fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border
                
            current_row += 1

        for col in ws_pol.columns:
            max_length = 0
            column_letter = col[0].column_letter
            for cell in col:
                try:
                    val = str(cell.value) if cell.value is not None else ""
                    lines = val.split("\n")
                    for line in lines:
                        if len(line) > max_length: max_length = len(line)
                except: pass
            ws_pol.column_dimensions[column_letter].width = min(max(max_length + 4, 15), 60)

        pol_max_row = ws_pol.max_row
        for col_idx in range(1, 7):
            start_merge_row = 2
            for row in range(2, pol_max_row + 2):
                if row <= pol_max_row:
                    current_tuple = tuple(policy_data[row-2].get(pol_headers[c-1], "") for c in range(1, col_idx + 1))
                else:
                    current_tuple = None
                    
                prev_tuple = tuple(policy_data[start_merge_row-2].get(pol_headers[c-1], "") for c in range(1, col_idx + 1))
                
                if current_tuple != prev_tuple or row > pol_max_row:
                    if row - 1 > start_merge_row and prev_tuple[-1] != "":
                        ws_pol.merge_cells(start_row=start_merge_row, start_column=col_idx, end_row=row-1, end_column=col_idx)
                    start_merge_row = row
    
    if datasets_data:
        ds_header_fill = PatternFill(start_color="C65911", end_color="C65911", fill_type="solid")
        
        for ds_name, ip_list in datasets_data.items():
            sheet_title = clean_sheet_name(f"DS_{ds_name}")
            ws_ds = wb.create_sheet(title=sheet_title)
            
            ws_ds.append(["ID", "Index (NLB)", "IP Adresi / Aralığı"])
            
            for cell in ws_ds[1]:
                cell.font = header_font
                cell.fill = ds_header_fill
                cell.alignment = center_alignment
                cell.border = thin_border
            
            for ip_data in ip_list:
                ws_ds.append([ip_data["ID"], ip_data["Index"], ip_data["IP"]])
                
            for row in ws_ds.iter_rows(min_row=2, max_row=ws_ds.max_row, min_col=1, max_col=3):
                for cell in row:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin_border
            
            ws_ds.column_dimensions["A"].width = 10
            ws_ds.column_dimensions["B"].width = 15
            ws_ds.column_dimensions["C"].width = 40

    wb.save(excel_path)