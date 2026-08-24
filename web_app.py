import os
import sys
import threading
import html
from flask import Flask, render_template, request, jsonify, send_file
import config_manager
import openpyxl
from openpyxl.utils.cell import range_boundaries
import main as core_main
import azure_manager

app = Flask(__name__, template_folder='web')

class LogCapture:
    def __init__(self):
        self.logs = []
        self.original_stdout = sys.stdout
        self.is_running = False

    def write(self, text):
        self.logs.append(text)
        self.original_stdout.write(text)

    def flush(self):
        self.original_stdout.flush()

logger = LogCapture()

def background_task(companies):
    sys.stdout = logger
    logger.is_running = True
    logger.logs = [] 
    try:
        core_main.run_reports(companies)
    except Exception as e:
        print(f"\n[!] Kritik Hata: {e}")
    finally:
        sys.stdout = logger.original_stdout
        logger.is_running = False

def get_css_color(pxl_color):
    if not pxl_color or getattr(pxl_color, 'type', None) != 'rgb' or not pxl_color.rgb:
        return None
    rgb = str(pxl_color.rgb)
    if rgb == '00000000': 
        return None
    if len(rgb) == 8:
        return f"#{rgb[2:]}"
    elif len(rgb) == 6:
        return f"#{rgb}"
    return None


@app.route('/')
def index():
    sirketler = config_manager.load_settings()
    companies = list(sirketler.keys())
    
    reports = {}
    if os.path.exists('raporlar'):
        files = sorted(os.listdir('raporlar'), reverse=True)
        for f in files:
            if f.endswith('.xlsx'):
                for c in companies:
                    clean_c = c.lower().replace(" ", "_")
                    if f.startswith(clean_c):
                        if c not in reports: reports[c] = []
                        reports[c].append(f)
                        break
                        
    return render_template('index.html', companies=companies, reports=reports)

@app.route('/run', methods=['POST'])
def run_report():
    if logger.is_running:
        return jsonify({"status": "error", "message": "Zaten çalışan bir işlem var! Lütfen bitmesini bekleyin."})
    
    companies = request.json.get('companies', [])
    if not companies:
        return jsonify({"status": "error", "message": "Şirket seçilmedi!"})
        
    threading.Thread(target=background_task, args=(companies,)).start()
    return jsonify({"status": "success"})

@app.route('/logs')
def get_logs():
    return jsonify({"logs": "".join(logger.logs), "is_running": logger.is_running})

@app.route('/download/<filename>')
def download(filename):
    filepath = os.path.join('raporlar', filename)
    if os.path.exists(filepath):
        return send_file(filepath, as_attachment=True)
    return "Dosya bulunamadı", 404

@app.route('/view/<filename>')
def view_excel(filename):
    filepath = os.path.join('raporlar', filename)
    if not os.path.exists(filepath):
        return "Dosya bulunamadı", 404
        
    wb = openpyxl.load_workbook(filepath, data_only=True)
    
    sheets_html = {}
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        tables_html = "<table class='excel-table'>"
        
        merged_cells_map = {}
        skip_cells = set()
        
        for merged_range in ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
            rowspan = max_row - min_row + 1
            colspan = max_col - min_col + 1
            
            top_left = (min_row, min_col)
            merged_cells_map[top_left] = {'rowspan': rowspan, 'colspan': colspan}
            
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    if (r, c) != top_left:
                        skip_cells.add((r, c))

        for r_idx, row in enumerate(ws.iter_rows(), start=1):
            tables_html += "<tr>"
            for c_idx, cell in enumerate(row, start=1):
                coord = (r_idx, c_idx)
                
                if coord in skip_cells:
                    continue
                    
                html_attrs = []
                css_styles = []
                
                if coord in merged_cells_map:
                    span_info = merged_cells_map[coord]
                    if span_info['rowspan'] > 1: html_attrs.append(f"rowspan='{span_info['rowspan']}'")
                    if span_info['colspan'] > 1: html_attrs.append(f"colspan='{span_info['colspan']}'")

                if cell.fill and cell.fill.start_color:
                    bg_hex = get_css_color(cell.fill.start_color)
                    if bg_hex: css_styles.append(f"background-color: {bg_hex} !important;")

                if cell.font:
                    if cell.font.color:
                        font_hex = get_css_color(cell.font.color)
                        if font_hex: css_styles.append(f"color: {font_hex} !important;")
                    if cell.font.bold:
                        css_styles.append("font-weight: bold;")
                        
                if cell.alignment:
                    if cell.alignment.horizontal:
                        css_styles.append(f"text-align: {cell.alignment.horizontal};")
                    if cell.alignment.vertical:
                        v_align = cell.alignment.vertical
                        if v_align == 'center': v_align = 'middle'
                        css_styles.append(f"vertical-align: {v_align};")
                    if cell.alignment.wrap_text:
                        css_styles.append("white-space: pre-wrap;")
                    else:
                        css_styles.append("white-space: nowrap;")

                style_str = f" style='{'; '.join(css_styles)}'" if css_styles else ""
                attr_str = f" {' '.join(html_attrs)}" if html_attrs else ""

                val = cell.value if cell.value is not None else ""
                val = html.escape(str(val))
                val = val.replace('\n', '<br>')
                
                tables_html += f"<td{attr_str}{style_str}>{val}</td>"
                
            tables_html += "</tr>\n"
            
        tables_html += "</table>"
        sheets_html[sheet_name] = tables_html
        
    return render_template('excel_web_view.html', sheets=sheets_html, filename=filename)

@app.route('/azure_test')
def azure_test_page():
    sirketler = config_manager.load_settings()
    azure_companies = [c for c, cfg in sirketler.items() if cfg.get("AZURE")]
    return render_template('azure_test.html', companies=azure_companies)

@app.route('/api/azure_targets', methods=['POST'])
def api_azure_targets():
    company = request.json.get('company')
    if not company:
        return jsonify({"status": "error", "message": "Şirket seçilmedi."})

    sirketler = config_manager.load_settings()
    sirket_config = sirketler.get(company, {})
    
    if not sirket_config.get("AZURE"):
        return jsonify({"status": "error", "message": "Bu şirket için Azure aktif değil."})

    az_url = os.getenv(f"{company}_AZURE_SERVER_URL")
    az_col = os.getenv(f"{company}_AZURE_COLLECTION")
    az_proj = os.getenv(f"{company}_AZURE_PROJECT")
    az_dg_id = os.getenv(f"{company}_AZURE_DEPLOYMENT_GROUP_ID")
    az_pat = os.getenv(f"{company}_AZURE_PAT")
    
    if not all([az_url, az_col, az_proj, az_dg_id, az_pat]):
        return jsonify({"status": "error", "message": "Eksik .env parametreleri. (URL, Proje, Token vb.)"})

    old_stdout = sys.stdout
    devnull = open(os.devnull, 'w', encoding='utf-8')
    sys.stdout = devnull
    
    try:
        targets = azure_manager.fetch_azure_deployment_targets(az_url, az_col, az_proj, az_dg_id, az_pat)
    except Exception as e:
        sys.stdout = old_stdout
        devnull.close()
        return jsonify({"status": "error", "message": str(e)})
    
    sys.stdout = old_stdout
    devnull.close()
    return jsonify({"status": "success", "data": targets})

def start_server(port):
    print(f"[*] http://127.0.0.1:{port} adresi üzerinde web arayüzü başlatıldı.\n")
    app.run(host='0.0.0.0', port=port, debug=False)